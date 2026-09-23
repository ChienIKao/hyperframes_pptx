#!/usr/bin/env python3
"""Markdown → PPTX builder for the md-to-pptx skill.

Subcommands
  build   deck.md -o out.pptx [--template SLUG|PATH] [--theme SLUG|PATH]
  inspect template.pptx [--write-config config.json]
  dump    deck.pptx            (content + overflow QA report)

Heading hierarchy (overridable in front matter):
  #    → cover
  ##   → chapter: outline entry + divider with sub-TOC   section_level: 2
  ###  → subsection: sub-TOC entry + nav bar 2nd row      subsection_level: 3
  #### → content slide                                    slide_level: 4
  `標題 | Subtitle` on any heading adds a subtitle.

Themes with "engine": "canvas" (default: blockframe) are drawn by canvas_deck.py;
.pptx/.potx templates use the placeholder renderer in this file.
"""
from __future__ import annotations

import argparse
import io
import json
import math
import re
import sys
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SKILL_DIR = Path(__file__).resolve().parent.parent
EMU_PER_PT = 12700

# --------------------------------------------------------------------------- model


@dataclass
class Block:
    kind: str  # bullets | para | sub | quote | image | table | code
    data: object


@dataclass
class Container:
    title: str = ""
    subtitle: str = ""
    blocks: list = field(default_factory=list)
    notes: list = field(default_factory=list)


@dataclass
class Subsection(Container):
    slides: list = field(default_factory=list)


@dataclass
class Section(Container):
    slides: list = field(default_factory=list)        # slides before the first subsection
    subsections: list = field(default_factory=list)


@dataclass
class Doc:
    meta: dict
    cover: Container
    sections: list


@dataclass
class SlideSpec:
    role: str  # cover | agenda | section | content | closing
    title: str = ""
    subtitle: str = ""
    blocks: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    label: str = ""


# --------------------------------------------------------------------------- markdown

HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
LIST_ITEM = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
IMAGE_LINE = re.compile(r'^!\[([^\]]*)\]\(\s*<?([^)\s>]+)>?(?:\s+"[^"]*")?\s*\)$')
FENCE = re.compile(r"^(`{3,}|~{3,})\s*([\w+-]*)\s*(.*)$")
HR = re.compile(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$")
TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
ALERT = re.compile(r"^\[!(\w+)\]\s*(.*)$")
SOURCE = re.compile(r"^(資料來源|来源|來源|參考資料|Sources?|Ref(?:erences?)?)\s*[:：]", re.I)
COMPONENT_TYPES = {"cards", "steps", "timeline", "flow", "stats", "compare", "chart"}
ATTR = re.compile(r"^(tag|img|icon|color)\s*[:：]\s*(.+)$", re.I)


def is_cjk(ch: str) -> bool:
    return unicodedata.east_asian_width(ch) in ("W", "F")


def join_lines(a: str, b: str) -> str:
    if not a:
        return b
    if a and b and (is_cjk(a[-1]) or is_cjk(b[0])):
        return a + b
    return a + " " + b


def split_frontmatter(text: str):
    meta = {}
    if text.startswith("---"):
        end = text.find("\n---", 3)
        if end != -1:
            for line in text[3:end].strip().splitlines():
                if ":" in line and not line.lstrip().startswith("#"):
                    k, v = line.split(":", 1)
                    v = v.strip().strip('"').strip("'")
                    if v.lower() in ("true", "false"):
                        v = v.lower() == "true"
                    meta[k.strip()] = v
            text = text[end + 4:]
    return meta, text


def split_row(line: str):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", s)]


def split_heading(text: str):
    """`標題 | Subtitle` → (標題, Subtitle)"""
    if " | " in text:
        a, b = text.split(" | ", 1)
        return a.strip(), b.strip()
    return text.strip(), ""


def make_item(text: str):
    pos, attrs = [], {}
    for part in re.split(r"\s\|\s", text):
        part = part.strip()
        m = ATTR.match(part)
        if m:
            attrs[m.group(1).lower()] = m.group(2).strip()
        elif part:
            pos.append(part)
    return {"title": pos[0] if pos else "", "desc": pos[1] if len(pos) > 1 else "",
            "extra": pos[2:], "attrs": attrs, "children": []}


def parse_component(kind: str, info: str, body: str):
    items, free = [], []
    for line in body.split("\n"):
        m = LIST_ITEM.match(line.replace("\t", "    "))
        if m and (not m.group(1) or not items):
            items.append(make_item(m.group(3)))
        elif m:
            items[-1]["children"].append(m.group(3).strip())
        elif line.strip() and items and line.startswith((" ", "\t")):
            items[-1]["children"].append(line.strip())
        elif line.strip():
            free.append(line.strip())
    if kind == "flow" and not items:
        for ln in free:
            if "->" in ln or "→" in ln:
                items += [make_item(x) for x in re.split(r"\s*(?:->|→)\s*", ln) if x.strip()]
    rows = [split_row(l) for l in free if l.startswith("|") and not TABLE_SEP.match(l)]
    return {"type": kind, "args": info.split(), "items": items, "lines": free, "rows": rows}


def parse_markdown(text: str) -> Doc:
    meta, body = split_frontmatter(text.replace("\r\n", "\n"))
    sec_lv = int(meta.get("section_level", 2))
    sub_lv = int(meta.get("subsection_level", 3))
    sl_lv = int(meta.get("slide_level", 4))
    doc = Doc(meta=meta, cover=Container(title=meta.get("title", "")), sections=[])
    cur: Container = doc.cover
    cur_section: Section | None = None
    cur_sub: Subsection | None = None
    title_seen = bool(meta.get("title"))

    para: list[str] = []
    bullets: list | None = None
    indents: list[int] = []

    def flush():
        nonlocal para, bullets, indents
        if para:
            text_ = ""
            for ln in para:
                text_ = join_lines(text_, ln.strip())
            cur.blocks.append(Block("source" if SOURCE.match(text_) else "para", text_))
            para = []
        if bullets:
            cur.blocks.append(Block("bullets", bullets))
        bullets, indents = None, []

    def ensure_section():
        nonlocal cur_section
        if cur_section is None:
            cur_section = Section(title="")
            doc.sections.append(cur_section)
        return cur_section

    lines = body.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # speaker notes: <!-- ... -->
        if stripped.startswith("<!--"):
            flush()
            chunk = [stripped[4:]]
            while "-->" not in chunk[-1] and i + 1 < len(lines):
                i += 1
                chunk.append(lines[i])
            note = "\n".join(chunk).split("-->")[0].strip()
            note = re.sub(r"^(notes?|備註|讲稿|講稿)\s*[:：]\s*", "", note, flags=re.I)
            if note:
                cur.notes.append(note)
            i += 1
            continue

        fm = FENCE.match(stripped)
        if fm:
            flush()
            fence, lang, info = fm.group(1), fm.group(2), fm.group(3)
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(fence):
                code.append(lines[i])
                i += 1
            if lang.lower() in COMPONENT_TYPES:
                cur.blocks.append(Block("component", parse_component(lang.lower(), info, "\n".join(code))))
            else:
                cur.blocks.append(Block("code", {"lang": lang, "text": "\n".join(code)}))
            i += 1
            continue

        hm = HEADING.match(line)
        if hm:
            flush()
            lvl = len(hm.group(1))
            title, subtitle = split_heading(hm.group(2))
            if lvl == 1 and not title_seen:
                doc.cover.title, doc.cover.subtitle = title, subtitle
                cur = doc.cover
                title_seen = True
            elif sec_lv and lvl <= sec_lv:
                cur_section = Section(title=title, subtitle=subtitle)
                doc.sections.append(cur_section)
                cur, cur_sub = cur_section, None
            elif sub_lv and lvl == sub_lv:
                ensure_section()
                cur_sub = Subsection(title=title, subtitle=subtitle)
                cur_section.subsections.append(cur_sub)
                cur = cur_sub
            elif lvl <= sl_lv:
                ensure_section()
                cur = Container(title=title, subtitle=subtitle)
                (cur_sub.slides if cur_sub else cur_section.slides).append(cur)
            else:
                cur.blocks.append(Block("sub", hm.group(2).strip()))
            i += 1
            continue

        if HR.match(line):
            flush()
            owner = cur_sub.slides if cur_sub else (cur_section.slides if cur_section else None)
            if owner is not None and (cur in owner or cur is cur_sub):
                cur = Container(title=cur.title, subtitle=cur.subtitle)
                owner.append(cur)
            i += 1
            continue

        im = IMAGE_LINE.match(stripped)
        if im:
            flush()
            cur.blocks.append(Block("image", {"alt": im.group(1), "path": im.group(2)}))
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1]):
            flush()
            header = split_row(lines[i])
            aligns = []
            for c in split_row(lines[i + 1]):
                aligns.append("r" if c.endswith(":") and not c.startswith(":")
                              else "c" if c.startswith(":") and c.endswith(":") else "l")
            rows = [header]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            cur.blocks.append(Block("table", {"rows": rows, "aligns": aligns}))
            continue

        if stripped.startswith(">"):
            flush()
            q = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q.append(lines[i].strip()[1:].strip())
                i += 1
            am = ALERT.match(q[0]) if q else None
            if am:
                q[0] = am.group(2)
            text_ = ""
            for ln in q:
                text_ = join_lines(text_, ln)
            if am:
                cur.blocks.append(Block("callout", {"kind": am.group(1).lower(), "text": text_}))
            else:
                cur.blocks.append(Block("quote", text_))
            continue

        lm = LIST_ITEM.match(line.replace("\t", "    "))
        if lm:
            if para:
                flush()
            indent = len(lm.group(1))
            if bullets is None:
                bullets, indents = [], []
            while indents and indent < indents[-1]:
                indents.pop()
            if not indents or indent > indents[-1]:
                indents.append(indent)
            level = min(len(indents) - 1, 4)
            ordered = lm.group(2)[0].isdigit()
            bullets.append({"level": level, "text": lm.group(3).strip(), "ordered": ordered})
            i += 1
            continue

        if not stripped:
            if para:
                flush()
            i += 1
            continue

        # continuation line of a list item, or plain paragraph text
        if bullets and line.startswith((" ", "\t")):
            bullets[-1]["text"] = join_lines(bullets[-1]["text"], stripped)
        else:
            if bullets:
                flush()
            para.append(stripped)
        i += 1

    flush()
    return doc


# --------------------------------------------------------------------------- inline

INLINE = re.compile(
    r"==(?P<hl>.+?)==|\*\*(?P<b>.+?)\*\*|__(?P<b2>.+?)__|`(?P<code>[^`]+)`|\[(?P<lt>[^\]]+)\]\((?P<lu>[^)]+)\)"
    r"|(?<![\w*])\*(?!\s)(?P<i>.+?)(?<!\s)\*(?!\*)|(?<!\w)_(?!\s)(?P<i2>.+?)(?<!\s)_(?!\w)"
)


def inline_runs(text: str):
    runs, pos = [], 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            runs.append((text[pos:m.start()], {}))
        if m.group("hl"):
            runs.append((m.group("hl"), {"hl": True, "bold": True}))
        elif m.group("b") or m.group("b2"):
            runs.append((m.group("b") or m.group("b2"), {"bold": True}))
        elif m.group("code"):
            runs.append((m.group("code"), {"code": True}))
        elif m.group("lt"):
            runs.append((m.group("lt"), {"link": m.group("lu")}))
        else:
            runs.append((m.group("i") or m.group("i2"), {"italic": True}))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], {}))
    return runs or [("", {})]


def plain(text: str) -> str:
    return "".join(t for t, _ in inline_runs(text))


def disp_width(s: str) -> float:
    """Approximate width in em."""
    w = 0.0
    for ch in s:
        if is_cjk(ch):
            w += 1.0
        elif ch == " ":
            w += 0.28
        elif ch.isupper() or ch.isdigit():
            w += 0.62
        else:
            w += 0.5
    return w


# --------------------------------------------------------------------------- planning


def text_units(block: Block):
    """Split a text block into atomic units (for pagination) with a weight each."""
    if block.kind == "bullets":
        groups, cur = [], []
        for it in block.data:
            if it["level"] == 0 and cur:
                groups.append(cur)
                cur = []
            cur.append(it)
        if cur:
            groups.append(cur)
        return [(Block("bullets", g), sum(1 + disp_width(plain(it["text"])) // 40 for it in g)) for g in groups]
    if block.kind in ("para", "quote"):
        return [(block, 1 + disp_width(plain(block.data)) // 40)]
    return [(block, 1)]


TEXT_KINDS = ("bullets", "para", "sub", "quote")


def merge_bullets(bs):
    """Merge adjacent bullet groups back into one block."""
    out = []
    for b in bs:
        if out and b.kind == "bullets" and out[-1].kind == "bullets":
            out[-1] = Block("bullets", out[-1].data + b.data)
        else:
            out.append(b)
    return out


def paginate(title, blocks, notes, cfg):
    max_units = int(cfg.get("max_lines", 9))
    max_rows = int(cfg.get("max_table_rows", 10))
    visuals, units = [], []
    for b in blocks:
        if b.kind == "table" and len(b.data["rows"]) - 1 > max_rows:
            head, body = b.data["rows"][0], b.data["rows"][1:]
            for k in range(0, len(body), max_rows):
                visuals.append(Block("table", {"rows": [head] + body[k:k + max_rows], "aligns": b.data["aligns"]}))
        elif b.kind in ("image", "table", "code"):
            visuals.append(b)
        else:
            units.extend(text_units(b))

    # with a visual beside the text, the text column holds less
    limit = max_units if not visuals else max(4, int(max_units * 0.75))
    pages, cur, w = [], [], 0
    for blk, wt in units:
        if cur and w + wt > limit:
            pages.append(cur)
            cur, w = [], 0
        cur.append(blk)
        w += wt
    if cur or not pages:
        pages.append(cur)

    specs = []
    extra_visuals = visuals[2:] if len(pages) == 1 and len(visuals) > 2 and units else []
    first_visuals = visuals[:2] if extra_visuals else visuals
    for n, pg in enumerate(pages):
        vis = first_visuals if n == 0 else []
        # tables beyond the first get their own slides
        tables = [v for v in vis if v.kind == "table"]
        if len(tables) > 1:
            vis = [v for v in vis if v.kind != "table"] + tables[:1]
            extra_visuals = tables[1:] + extra_visuals
        specs.append(SlideSpec("content", title=title, blocks=merge_bullets(pg) + vis, notes=notes if n == 0 else []))
    for v in extra_visuals:
        specs.append(SlideSpec("content", title=title, blocks=[v]))
    cont = cfg.get("continued_suffix", "（續）")
    for s in specs[1:]:
        s.title = f"{title}{cont}" if title else ""
    return specs


def degrade(blocks):
    """Map canvas-only blocks onto what the template engine can draw."""
    out = []
    for b in blocks:
        if b.kind == "component":
            d = b.data
            if d["type"] == "chart" and d["rows"]:
                out.append(Block("table", {"rows": d["rows"], "aligns": []}))
                continue
            items = []
            for it in d["items"]:
                t = f"**{it['title']}**" + (f"：{it['desc']}" if it["desc"] else "")
                items.append({"level": 0, "text": t, "ordered": d["type"] in ("steps", "timeline", "flow")})
                items += [{"level": 1, "text": c, "ordered": False} for c in it["children"]]
            if items:
                out.append(Block("bullets", items))
        elif b.kind == "callout":
            out.append(Block("quote", b.data["text"]))
        elif b.kind == "source":
            out.append(Block("para", b.data))
        else:
            out.append(b)
    return out


def plan(doc: Doc, cfg: dict):
    meta = doc.meta
    slides = []
    cover = doc.cover
    subtitle = meta.get("subtitle", "")
    intro = cover.blocks
    if not subtitle and len(intro) == 1 and intro[0].kind in ("para", "quote") and disp_width(intro[0].data) <= 60:
        subtitle, intro = intro[0].data, []
    extra = " · ".join(str(meta[k]) for k in ("author", "date") if meta.get(k))
    slides.append(SlideSpec("cover", title=cover.title, subtitle="\n".join(x for x in (subtitle, extra) if x),
                            notes=cover.notes))
    if intro:
        slides += paginate(cover.title, intro, [], cfg)

    named = [s for s in doc.sections if s.title]
    if meta.get("agenda", cfg.get("agenda", True)) and len(named) >= 2:
        slides.append(SlideSpec("agenda", title=meta.get("agenda_title", cfg.get("agenda_title", "目錄")),
                                blocks=[Block("agenda", [s.title for s in named])]))

    label_fmt = meta.get("section_label", cfg.get("section_label", "{n:02d}"))
    n = 0
    for sec in doc.sections:
        if sec.title:
            n += 1
            sub, body = "", sec.blocks
            if len(body) == 1 and body[0].kind in ("para", "quote") and disp_width(body[0].data) <= 60:
                sub, body = body[0].data, []
            slides.append(SlideSpec("section", title=sec.title, subtitle=sub, notes=sec.notes,
                                    label=label_fmt.format(n=n)))
            if body:
                slides += paginate(sec.title, degrade(body), [], cfg)
        elif sec.blocks:
            slides += paginate("", degrade(sec.blocks), sec.notes, cfg)
        for s in sec.slides:
            slides += paginate(s.title, degrade(s.blocks), s.notes, cfg)
        for sub in sec.subsections:
            if sub.blocks or not sub.slides:
                slides += paginate(sub.title, degrade(sub.blocks), sub.notes, cfg)
            for s in sub.slides:
                slides += paginate(s.title, degrade(s.blocks), s.notes, cfg)

    closing = meta.get("closing", cfg.get("closing", False))
    if closing:
        slides.append(SlideSpec("closing", title=str(closing) if closing is not True else "謝謝聆聽"))
    return slides


# --------------------------------------------------------------------------- pptx helpers

TEMPLATE_CT = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
PRES_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"


def open_presentation(path: Path | None):
    if path is None:
        return Presentation()
    data = path.read_bytes()
    if path.suffix.lower() == ".potx":
        src = zipfile.ZipFile(io.BytesIO(data))
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                content = src.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    content = content.replace(TEMPLATE_CT.encode(), PRES_CT.encode())
                dst.writestr(item, content)
        data = buf.getvalue()
    return Presentation(io.BytesIO(data))


def remove_all_slides(prs):
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)


def hex_rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


def is_dark(h: str | None) -> bool:
    if not h:
        return False
    r, g, b = (int(h.lstrip("#")[k:k + 2], 16) for k in (0, 2, 4))
    return 0.299 * r + 0.587 * g + 0.114 * b < 128


def set_font(run, family=None, ea=None, size=None, color=None, bold=None, italic=None):
    f = run.font
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color:
        f.color.rgb = hex_rgb(color)
    if family:
        f.name = family
    if ea:
        rPr = run._r.get_or_add_rPr()
        el = rPr.find(qn("a:ea"))
        if el is None:
            el = etree.SubElement(rPr, qn("a:ea"))
            latin = rPr.find(qn("a:latin"))
            anchor = None
            for tag in ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"):
                anchor = rPr.find(qn(tag))
                if anchor is not None:
                    break
            rPr.remove(el)
            if latin is not None:
                latin.addnext(el)
            elif anchor is not None:
                anchor.addprevious(el)
            else:
                rPr.append(el)
        el.set("typeface", ea)


BU_TAGS = ("a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip")
PPR_TAIL = ("a:tabLst", "a:defRPr", "a:extLst")


def set_bullet(p, mode, level=0):
    """mode: inherit | none | num | char"""
    if mode == "inherit":
        return
    pPr = p._p.get_or_add_pPr()
    for tag in BU_TAGS:
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if mode == "none":
        el = etree.Element(qn("a:buNone"))
        pPr.set("marL", "0")
        pPr.set("indent", "0")
    elif mode == "num":
        el = etree.Element(qn("a:buAutoNum"), type="arabicPeriod")
    else:
        el = etree.Element(qn("a:buChar"), char="•")
        pPr.set("marL", str(int(Inches(0.3) * (level + 1))))
        pPr.set("indent", str(-int(Inches(0.25))))
    anchor = None
    for tag in PPR_TAIL:
        anchor = pPr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(el)
    else:
        pPr.append(el)


def box(shape):
    return (shape.left, shape.top, shape.width, shape.height)


def estimate_height_pt(paras, size, width_emu):
    width_pt = width_emu / EMU_PER_PT - 14.4
    total = 0.0
    for p in paras:
        s = size * p.get("scale", 1.0)
        avail = max(40.0, width_pt - 22 * p.get("level", 0) - (18 if p.get("bullet") != "none" else 0))
        lines = 0
        for seg in plain(p["text"]).split("\n"):
            lines += max(1, math.ceil(disp_width(seg) * s / avail))
        total += lines * s * 1.2 + s * 0.45
    return total + 7.2


def fit_size(paras, w, h, max_size, min_size):
    for s in range(int(max_size), int(min_size) - 1, -1):
        if estimate_height_pt(paras, s, w) <= h / EMU_PER_PT:
            return s, True
    return int(min_size), False


# --------------------------------------------------------------------------- renderer

ROLE_KEYWORDS = {
    "cover": ["title slide", "標題投影片", "标题幻灯片", "封面", "cover"],
    "section": ["section", "章節", "章节", "節標題", "节标题", "區段", "divider"],
    "content": ["title and content", "標題及內容", "标题和内容", "標題及物件", "content"],
    "two_content": ["two content", "兩個內容", "两栏内容", "两项内容", "two column"],
    "title_only": ["title only", "只有標題", "仅标题", "僅標題"],
    "closing": ["closing", "end", "結尾", "结束", "thank"],
}
BODY_TYPES = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
META_TYPES = (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)


def ph_types(layout):
    return [ph.placeholder_format.type for ph in layout.placeholders]


def detect_layouts(prs, mapping: dict):
    layouts = list(prs.slide_layouts)
    by_name = {l.name: l for l in layouts}
    found = {}
    for role, ref in (mapping or {}).items():
        if isinstance(ref, int) or (isinstance(ref, str) and ref.isdigit()):
            found[role] = layouts[int(ref)]
        elif ref in by_name:
            found[role] = by_name[ref]
        else:
            print(f"warning: layout '{ref}' for role '{role}' not found", file=sys.stderr)

    def structural(role, l):
        t = ph_types(l)
        bodies = sum(1 for x in t if x in BODY_TYPES)
        has_title = any(x in TITLE_TYPES for x in t)
        return {
            "cover": PP_PLACEHOLDER.CENTER_TITLE in t,
            "content": has_title and bodies == 1,
            "two_content": has_title and bodies == 2,
            "title_only": has_title and bodies == 0,
            "section": has_title and bodies == 1,
        }.get(role, False)

    for role, kws in ROLE_KEYWORDS.items():
        if role in found:
            continue
        for l in layouts:
            if any(k in l.name.lower() for k in kws) and (role == "closing" or structural(role, l) or role == "section"):
                found[role] = l
                break
        else:
            for l in layouts:
                if role != "closing" and structural(role, l):
                    found[role] = l
                    break
    found.setdefault("content", layouts[min(1, len(layouts) - 1)])
    found.setdefault("cover", layouts[0])
    found.setdefault("section", found["cover"])
    found.setdefault("title_only", found["content"])
    found.setdefault("agenda", found["content"])
    found.setdefault("closing", found["section"])
    return found


def scale_to_widescreen(prs):
    """Default python-pptx template is 4:3; stretch every master/layout placeholder to 16:9."""
    ratio = Inches(13.333) / prs.slide_width
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    master = prs.slide_master
    for shp in list(master.shapes) + [s for l in prs.slide_layouts for s in l.shapes]:
        # only shapes with their own xfrm; inherited geometry follows the (already scaled) master
        for xfrm in shp._element.xpath("./p:spPr/a:xfrm"):
            off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
            off.set("x", str(int(int(off.get("x")) * ratio)))
            ext.set("cx", str(int(int(ext.get("cx")) * ratio)))


class Renderer:
    def __init__(self, prs, cfg, theme, base_dir: Path):
        self.prs, self.cfg, self.theme, self.base = prs, cfg, theme, base_dir
        self.layouts = detect_layouts(prs, cfg.get("layouts", {}))
        self.warnings = []
        t = theme or {}
        self.fonts = t.get("fonts", {})
        self.colors = t.get("colors", {})
        self.bgs = t.get("backgrounds", {})
        self.sizes = {"title": None, "cover_title": None, "body_max": 24, "body_min": 14,
                      "code_max": 20, "code_min": 10, "table_max": 20, "table_min": 10, "caption": 12}
        self.sizes.update(cfg.get("sizes", {}))
        self.sizes.update(t.get("sizes", {}))
        self.title_align = t.get("title_align") or cfg.get("title_align")

    # ---- styling
    def dark(self, role):
        return is_dark(self.bgs.get(role))

    def color(self, role, key):
        if not self.colors:
            return None
        if self.dark(role):
            return self.colors.get(f"{key}_on_dark") or self.colors.get("text_on_dark")
        return self.colors.get(key)

    def write_runs(self, p, text, role, size=None, bold=None, italic=None, color_key="text", heading=False):
        fam = self.fonts.get("heading" if heading else "body")
        ea = self.fonts.get("ea_heading" if heading else "ea_body") or self.fonts.get("ea")
        for chunk, fmt in inline_runs(text):
            if not chunk:
                continue
            r = p.add_run()
            r.text = chunk
            c = self.color(role, "accent") if fmt.get("bold") and self.colors.get("emphasis_accent") else self.color(role, color_key)
            if fmt.get("code"):
                set_font(r, family=self.fonts.get("code", "Consolas"), ea=ea, size=size, color=c)
            else:
                set_font(r, family=fam, ea=ea, size=size, color=c,
                         bold=True if fmt.get("bold") else bold, italic=True if fmt.get("italic") else italic)
            if fmt.get("link"):
                r.hyperlink.address = fmt["link"]

    def fill_paras(self, tf, paras, role, size, textbox=False):
        tf.word_wrap = True
        # drop every existing paragraph but the first, clear the first
        for extra in tf.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        first = tf.paragraphs[0]
        for r in list(first.runs):
            r._r.getparent().remove(r._r)
        for n, pd in enumerate(paras):
            p = first if n == 0 else tf.add_paragraph()
            p.level = pd.get("level", 0)
            mode = pd.get("bullet", "inherit")
            if textbox and mode == "inherit":
                mode = "char"
            set_bullet(p, mode, pd.get("level", 0))
            if pd.get("align"):
                p.alignment = pd["align"]
            s = max(10, round(size * pd.get("scale", 1.0)))
            p.space_after = Pt(round(s * 0.35))
            self.write_runs(p, pd["text"], role, size=s, bold=pd.get("bold"), italic=pd.get("italic"),
                            color_key=pd.get("color", "text"))

    def set_title(self, slide, text, role):
        ph = slide.shapes.title
        if ph is None:
            return
        tf = ph.text_frame
        tf.text = ""
        p = tf.paragraphs[0]
        size = self.sizes.get("cover_title") if role in ("cover", "closing") else self.sizes.get("title")
        if self.title_align and role in ("content", "agenda"):
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[self.title_align]
        self.write_runs(p, text, role, size=size, color_key="title", heading=True,
                        bold=(self.theme or {}).get("title_bold"))

    def background(self, slide, role):
        c = self.bgs.get(role) or (self.bgs.get("content") if role in ("agenda",) else None)
        if c:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = hex_rgb(c)

    def new_slide(self, role, layout_role=None):
        slide = self.prs.slides.add_slide(self.layouts[layout_role or role])
        self.background(slide, role)
        return slide

    def placeholders(self, slide, types):
        return sorted([ph for ph in slide.placeholders if ph.placeholder_format.type in types],
                      key=lambda p: (p.left or 0, p.top or 0))

    def cleanup(self, slide):
        for ph in list(slide.placeholders):
            if ph.placeholder_format.type in META_TYPES:
                continue
            if ph.has_text_frame and ph.text_frame.text.strip():
                continue
            ph._element.getparent().remove(ph._element)

    def notes(self, slide, notes):
        if notes:
            slide.notes_slide.notes_text_frame.text = "\n\n".join(notes)

    def content_area(self):
        lay = self.layouts["content"]
        bodies = [ph for ph in lay.placeholders if ph.placeholder_format.type in BODY_TYPES]
        if bodies and bodies[0].width:
            return box(bodies[0])
        W, H = self.prs.slide_width, self.prs.slide_height
        return (int(W * 0.06), int(H * 0.22), int(W * 0.88), int(H * 0.70))

    def split_to_fit(self, spec: SlideSpec):
        """Halve a content slide until its text fits the real box at `split_below` pt."""
        if spec.role != "content":
            return [spec]
        text = [b for b in spec.blocks if b.kind in TEXT_KINDS]
        vis = [b for b in spec.blocks if b.kind not in TEXT_KINDS]
        units = [u for b in text for u, _ in text_units(b)]
        if len(units) < 2:
            return [spec]
        _, _, w, h = self.content_area()
        if vis:
            w = int(w * float(self.cfg.get("text_ratio", 0.45)))
        floor = self.sizes.get("split_below", 16)
        if fit_size(self.block_paras(text), w, h, floor, floor)[1]:
            return [spec]
        mid = len(units) // 2
        suffix = self.cfg.get("continued_suffix", "（續）")
        cont = spec.title if spec.title.endswith(suffix) or not spec.title else spec.title + suffix
        a = SlideSpec("content", spec.title, blocks=merge_bullets(units[:mid]) + vis, notes=spec.notes)
        b = SlideSpec("content", cont, blocks=merge_bullets(units[mid:]))
        return self.split_to_fit(a) + self.split_to_fit(b)

    # ---- slide kinds
    def render(self, spec: SlideSpec):
        getattr(self, f"render_{spec.role}")(spec)

    def render_cover(self, spec):
        s = self.new_slide("cover")
        self.set_title(s, spec.title, "cover")
        subs = self.placeholders(s, (PP_PLACEHOLDER.SUBTITLE,) + BODY_TYPES)
        if subs and spec.subtitle:
            paras = [{"text": ln, "bullet": "none", "color": "muted" if k else "text"}
                     for k, ln in enumerate(spec.subtitle.split("\n"))]
            size, _ = fit_size(paras, subs[0].width, subs[0].height, self.sizes.get("subtitle", 22), 12)
            self.fill_paras(subs[0].text_frame, paras, "cover", size)
            for p in subs[0].text_frame.paragraphs:
                p.alignment = None  # inherit layout alignment
        self.notes(s, spec.notes)
        self.cleanup(s)

    def render_section(self, spec):
        s = self.new_slide("section")
        self.set_title(s, spec.title, "section")
        bodies = self.placeholders(s, BODY_TYPES + (PP_PLACEHOLDER.SUBTITLE,))
        text = "\n".join(x for x in (spec.label, spec.subtitle) if x)
        if bodies and text:
            paras = [{"text": ln, "bullet": "none", "color": "accent" if k == 0 and spec.label else "muted"}
                     for k, ln in enumerate(text.split("\n"))]
            self.fill_paras(bodies[0].text_frame, paras, "section", self.sizes.get("section_label", 24))
        self.notes(s, spec.notes)
        self.cleanup(s)

    def render_closing(self, spec):
        s = self.new_slide("closing")
        self.set_title(s, spec.title, "closing")
        self.cleanup(s)

    def render_agenda(self, spec):
        s = self.new_slide("agenda")
        self.set_title(s, spec.title, "agenda")
        items = spec.blocks[0].data
        paras = [{"text": f"{k + 1:02d}　{t}", "bullet": "none"} for k, t in enumerate(items)]
        bodies = self.placeholders(s, BODY_TYPES)
        if bodies:
            w, h = bodies[0].width, bodies[0].height
            size, _ = fit_size(paras, w, h, self.sizes["body_max"] + 8, self.sizes["body_min"])
            self.fill_paras(bodies[0].text_frame, paras, "agenda", size)
        self.cleanup(s)

    def block_paras(self, blocks):
        paras = []
        for b in blocks:
            if b.kind == "bullets":
                for it in b.data:
                    paras.append({"text": it["text"], "level": it["level"],
                                  "bullet": "num" if it["ordered"] else "inherit",
                                  "scale": max(0.7, 1 - 0.1 * it["level"])})
            elif b.kind == "para":
                paras.append({"text": b.data, "bullet": "none"})
            elif b.kind == "sub":
                paras.append({"text": b.data, "bullet": "none", "bold": True, "color": "accent", "scale": 1.05})
            elif b.kind == "quote":
                paras.append({"text": f"「{b.data}」", "bullet": "none", "italic": True, "color": "muted"})
        return paras

    def render_content(self, spec):
        text_blocks = [b for b in spec.blocks if b.kind in TEXT_KINDS]
        visuals = [b for b in spec.blocks if b.kind not in TEXT_KINDS]
        paras = self.block_paras(text_blocks)

        if not visuals:
            s = self.new_slide("content")
            self.set_title(s, spec.title, "content")
            bodies = self.placeholders(s, BODY_TYPES)
            if bodies and paras:
                self.fill_text(bodies[0], paras, spec.title)
        elif not paras:
            s = self.new_slide("content", "title_only")
            self.set_title(s, spec.title, "content")
            self.place_visuals(s, visuals, self.content_area(), spec.title)
        else:
            use_two = "two_content" in self.layouts and len(
                [p for p in self.layouts["two_content"].placeholders if p.placeholder_format.type in BODY_TYPES]) >= 2
            s = self.new_slide("content", "two_content" if use_two else "content")
            self.set_title(s, spec.title, "content")
            bodies = self.placeholders(s, BODY_TYPES)
            if use_two and len(bodies) >= 2:
                left, right = bodies[0], bodies[1]
                vbox = box(right)
                right._element.getparent().remove(right._element)
            else:
                left = bodies[0] if bodies else None
                x, y, w, h = self.content_area() if left is None else box(left)
                ratio = float(self.cfg.get("text_ratio", 0.45))
                gap = int(Inches(0.3))
                tw = int(w * ratio)
                vbox = (x + tw + gap, y, w - tw - gap, h)
                if left is not None:
                    left.left, left.top, left.width, left.height = x, y, tw, h
            if left is not None:
                self.fill_text(left, paras, spec.title)
            self.place_visuals(s, visuals, vbox, spec.title)
        self.notes(s, spec.notes)
        self.cleanup(s)

    def fill_text(self, ph, paras, title):
        size, ok = fit_size(paras, ph.width, ph.height, self.sizes["body_max"], self.sizes["body_min"])
        if not ok:
            self.warnings.append(f"「{title}」文字可能溢出（已縮到 {size}pt），建議拆頁或精簡")
        self.fill_paras(ph.text_frame, paras, "content", size)
        ph.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    def place_visuals(self, slide, visuals, area, title):
        x, y, w, h = area
        n = len(visuals)
        gap = int(Inches(0.25))
        if all(v.kind == "image" for v in visuals):
            cw = (w - gap * (n - 1)) // n
            cells = [(x + k * (cw + gap), y, cw, h) for k in range(n)]
        else:
            ch = (h - gap * (n - 1)) // n
            cells = [(x, y + k * (ch + gap), w, ch) for k in range(n)]
        for v, cell in zip(visuals, cells):
            getattr(self, f"place_{v.kind}")(slide, v.data, cell, title)

    def place_image(self, slide, d, cell, title):
        path = (self.base / d["path"]).resolve() if not Path(d["path"]).is_absolute() else Path(d["path"])
        x, y, w, h = cell
        if not path.exists():
            self.warnings.append(f"「{title}」找不到圖片：{d['path']}")
            return
        if path.suffix.lower() == ".svg":
            self.warnings.append(f"「{title}」SVG 不支援，請先轉成 PNG：{d['path']}")
            return
        cap_h = int(Inches(0.4)) if d["alt"] and self.cfg.get("image_captions", True) else 0
        from PIL import Image
        with Image.open(path) as im:
            iw, ih = im.size
        avail_h = h - cap_h
        scale = min(w / iw, avail_h / ih)
        pw, ph_ = int(iw * scale), int(ih * scale)
        px, py = x + (w - pw) // 2, y + (avail_h - ph_) // 2
        pic = slide.shapes.add_picture(str(path), px, py, pw, ph_)
        pic._element.nvPicPr.cNvPr.set("descr", d["alt"] or path.stem)
        if cap_h:
            tb = slide.shapes.add_textbox(x, py + ph_ + int(Inches(0.05)), w, cap_h)
            self.fill_paras(tb.text_frame, [{"text": d["alt"], "bullet": "none", "color": "muted",
                                             "align": PP_ALIGN.CENTER}], "content", self.sizes["caption"])

    def place_table(self, slide, d, cell, title):
        rows = d["rows"]
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        x, y, w, h = cell
        size = max(self.sizes["table_min"], min(self.sizes["table_max"], int(28 - 1.2 * len(rows) - ncols)))
        row_h = int(Pt(size * 2.1))
        shape = slide.shapes.add_table(len(rows), ncols, x, y, w, min(h, row_h * len(rows)))
        tbl = shape.table
        widths = [max(disp_width(plain(r[c])) for r in rows) + 2 for c in range(ncols)]
        tot = sum(widths)
        for c in range(ncols):
            tbl.columns[c].width = int(w * widths[c] / tot)
        aligns = d.get("aligns", [])
        amap = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
        for r, row in enumerate(rows):
            tbl.rows[r].height = row_h
            for c, val in enumerate(row):
                tf = tbl.cell(r, c).text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                if c < len(aligns):
                    p.alignment = amap[aligns[c]]
                fam = self.fonts.get("body")
                ea = self.fonts.get("ea_body") or self.fonts.get("ea")
                for chunk, fmt in inline_runs(val):
                    if chunk:
                        run = p.add_run()
                        run.text = chunk
                        set_font(run, family=fam, ea=ea, size=size, bold=True if (r == 0 or fmt.get("bold")) else None)
        est = row_h * len(rows)
        if est > h * 1.1:
            self.warnings.append(f"「{title}」表格可能超出版面（{len(rows)} 列）")

    def place_code(self, slide, d, cell, title):
        x, y, w, h = cell
        lines = d["text"].split("\n")
        longest = max((len(l) for l in lines), default=1)
        size = self.sizes["code_max"]
        while size > self.sizes["code_min"] and (
                longest * size * 0.6 > (w / EMU_PER_PT - 22) or len(lines) * size * 1.25 > h / EMU_PER_PT - 22):
            size -= 1
        tb = slide.shapes.add_textbox(x, y, w, min(h, int(Pt(len(lines) * size * 1.25 + 26))))
        fill = tb.fill
        fill.solid()
        fill.fore_color.rgb = hex_rgb(self.colors.get("code_bg", "F3F4F6"))
        tf = tb.text_frame
        tf.word_wrap = True
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Inches(0.15))
        for k, ln in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = ln if ln else " "
            set_font(r, family=self.fonts.get("code", "Consolas"), size=size,
                     color=self.colors.get("code_text", "1F2937"))
        if len(lines) * size * 1.25 > h / EMU_PER_PT:
            self.warnings.append(f"「{title}」程式碼區塊可能溢出（{len(lines)} 行）")


# --------------------------------------------------------------------------- config resolution


def load_json(p: Path):
    return json.loads(p.read_text(encoding="utf-8"))


def resolve_template(ref: str | None):
    """Returns (template_path|None, config dict)."""
    if not ref:
        return None, {}
    p = Path(ref)
    if p.is_dir() or (SKILL_DIR / "templates" / ref).is_dir():
        d = p if p.is_dir() else SKILL_DIR / "templates" / ref
        cfg = load_json(d / "config.json") if (d / "config.json").exists() else {}
        tpl = d / cfg.get("file", "")
        if not cfg.get("file") or not tpl.exists():
            cands = list(d.glob("*.pptx")) + list(d.glob("*.potx"))
            tpl = cands[0] if cands else None
        return tpl, cfg
    if p.exists():
        side = p.with_suffix(".json")
        return p, (load_json(side) if side.exists() else {})
    raise SystemExit(f"template not found: {ref}")


def resolve_theme(ref: str | None):
    if not ref:
        return None
    p = Path(ref)
    if not p.exists():
        p = SKILL_DIR / "themes" / f"{ref}.json"
    if not p.exists():
        raise SystemExit(f"theme not found: {ref}")
    return load_json(p)


# --------------------------------------------------------------------------- commands


def cmd_build(a):
    md_path = Path(a.markdown).resolve()
    doc = parse_markdown(md_path.read_text(encoding="utf-8"))
    tpl_ref = a.template or doc.meta.get("template")
    tpl_path, cfg = resolve_template(tpl_ref)
    theme_ref = a.theme or doc.meta.get("theme") or cfg.get("theme_ref") or (None if tpl_path else "blockframe")
    theme = resolve_theme(theme_ref) if isinstance(theme_ref, str) else None
    if cfg.get("theme"):
        theme = {**(theme or {}), **cfg["theme"]}

    if theme and theme.get("engine") == "canvas" and tpl_path is None:
        from canvas_deck import build_canvas
        base = Path(a.assets).resolve() if a.assets else md_path.parent
        prs, c = build_canvas(doc, theme, base)
        out = Path(a.output)
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(out)
        print(f"✔ {out}  ({len(prs.slides)} slides, engine=canvas, theme={theme.get('name', theme_ref)})")
        for w in c.warnings:
            print("  ⚠", w)
        return

    prs = open_presentation(tpl_path)
    if tpl_path is None:
        scale_to_widescreen(prs)
    else:
        remove_all_slides(prs)

    specs = plan(doc, {**cfg, **{k: v for k, v in doc.meta.items() if k in (
        "max_lines", "max_table_rows", "closing", "agenda", "agenda_title", "section_label")}})
    base = Path(a.assets).resolve() if a.assets else md_path.parent
    r = Renderer(prs, cfg, theme, base)
    count = 0
    for spec in specs:
        for piece in r.split_to_fit(spec):
            r.render(piece)
            count += 1

    out = Path(a.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"✔ {out}  ({count} slides, template={tpl_path.name if tpl_path else 'built-in'}, "
          f"theme={theme_ref if theme else '-'})")
    roles = {k: v.name for k, v in r.layouts.items()}
    print("  layouts:", json.dumps(roles, ensure_ascii=False))
    for w in r.warnings:
        print("  ⚠", w)


def cmd_inspect(a):
    prs = open_presentation(Path(a.template))
    W, H = prs.slide_width, prs.slide_height
    print(f"slide size: {W / 914400:.2f}in × {H / 914400:.2f}in   masters: {len(prs.slide_masters)}   "
          f"existing slides: {len(prs.slides)}")
    for i, l in enumerate(prs.slide_layouts):
        phs = []
        for ph in l.placeholders:
            t = str(ph.placeholder_format.type).split(".")[-1].split(" ")[0]
            if ph.placeholder_format.type in META_TYPES:
                continue
            geo = f"@{(ph.left or 0) / 914400:.1f},{(ph.top or 0) / 914400:.1f} {(ph.width or 0) / 914400:.1f}×{(ph.height or 0) / 914400:.1f}in"
            phs.append(f"{t}[{ph.placeholder_format.idx}]{geo}")
        print(f"[{i:2d}] {l.name}\n      " + ("  ".join(phs) or "(no placeholders)"))
    det = detect_layouts(prs, {})
    mapping = {k: v.name for k, v in det.items()}
    print("\nauto-detected roles:", json.dumps(mapping, ensure_ascii=False, indent=2))
    if a.write_config:
        cfg = {"file": Path(a.template).name, "layouts": mapping,
               "sizes": {"body_max": 24, "body_min": 14}, "theme": {"fonts": {}}}
        Path(a.write_config).write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\nconfig written → {a.write_config}")


def text_need(shp):
    """(needed, available) text height in pt for a shape's text frame, honouring its real margins."""
    tf = shp.text_frame
    w = (shp.width - tf.margin_left - tf.margin_right) / EMU_PER_PT
    h = (shp.height - tf.margin_top - tf.margin_bottom) / EMU_PER_PT
    need, last_sa = 0.0, 0.0
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        if not txt.strip():
            continue
        size = next((r.font.size.pt for r in p.runs if r.font.size), 18)
        last_sa = p.space_after.pt if p.space_after is not None else 0
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        pPr = p._p.pPr
        ind = int(pPr.get("marL", 0)) / EMU_PER_PT if pPr is not None else 0
        lines = sum(max(1, math.ceil(disp_width(seg) * size * 0.97 / max(20.0, w - ind))) for seg in txt.split("\n"))
        need += lines * size * 1.2 * ls + last_sa
    return need - last_sa, h


def cmd_dump(a):
    prs = Presentation(a.pptx)
    W, H = prs.slide_width, prs.slide_height
    tol = Inches(0.15)
    problems = 0
    for n, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text_frame.text if slide.shapes.title is not None else ""
        print(f"\n── {n:02d} [{slide.slide_layout.name}] {title}")
        for shp in slide.shapes:
            if shp.left is not None and (shp.left < -tol or shp.top < -tol or shp.left + shp.width > W + tol
                                         or shp.top + shp.height > H + tol):
                print(f"   ⚠ 超出投影片邊界：{shp.name}")
                problems += 1
            if shp.has_text_frame and shp != slide.shapes.title:
                for p in shp.text_frame.paragraphs:
                    txt = "".join(r.text for r in p.runs)
                    if txt.strip():
                        print(f"   {'  ' * p.level}· {txt}")
                need, avail = text_need(shp)
                if need > avail * 1.1 + 4:
                    print(f"   ⚠ 文字可能溢出（估計 {need:.0f}pt > 框高 {avail:.0f}pt）：{shp.text_frame.text[:20]}")
                    problems += 1
            elif shp.shape_type == 13:
                print(f"   [圖片] {shp._element.nvPicPr.cNvPr.get('descr', '')}")
            elif shp.has_chart:
                print(f"   [圖表] {shp.chart.chart_type}")
            elif shp.has_table:
                print(f"   [表格] {len(shp.table.rows)}×{len(shp.table.columns)}")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            print(f"   [備註] {slide.notes_slide.notes_text_frame.text.strip()[:60]}…")
    print(f"\n{len(prs.slides)} slides, {problems} potential problem(s)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build")
    b.add_argument("markdown")
    b.add_argument("-o", "--output", required=True)
    b.add_argument("--template", help="templates/<slug> or a .pptx/.potx path")
    b.add_argument("--theme", help="themes/<slug>.json or a path (default: blockframe)")
    b.add_argument("--assets", help="base dir for relative image paths (default: markdown's dir)")
    i = sub.add_parser("inspect")
    i.add_argument("template")
    i.add_argument("--write-config")
    d = sub.add_parser("dump")
    d.add_argument("pptx")
    a = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    {"build": cmd_build, "inspect": cmd_inspect, "dump": cmd_dump}[a.cmd](a)


if __name__ == "__main__":
    main()
