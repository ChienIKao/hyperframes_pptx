"""Canvas engine: draws every slide from design tokens (no template placeholders).

Used when the theme JSON has `"engine": "canvas"` (e.g. themes/blockframe.json).
Structure:  # cover · ## chapter · ### subsection · #### slide
Generated:  cover → outline (大綱) → per chapter: divider with sub-TOC → content slides
            (each with a top nav bar showing current chapter + subsection) → closing
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_PATTERN
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from build_deck import (TEXT_KINDS, Block, disp_width, estimate_height_pt, inline_runs, is_cjk, merge_bullets,
                        plain, text_units)

IN = 914400
W, H = 13.333, 7.5
PAD = 0.55
NAV_H = 1.0
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]


def E(x: float) -> int:
    return int(round(x * IN))


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


@dataclass
class Page:
    kind: str                      # cover | outline | divider | content | closing
    title: str = ""
    subtitle: str = ""
    chapter: int | None = None     # index into chapters
    sub: int | None = None         # index into chapter's subsections
    blocks: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    recap: bool = False


# --------------------------------------------------------------------------- XML helpers

RPR_AFTER_HL = ("a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill", "a:latin", "a:ea", "a:cs", "a:sym",
                "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")


def _insert_before(parent, el, tags):
    for t in tags:
        ref = parent.find(qn(t))
        if ref is not None:
            ref.addprevious(el)
            return
    parent.append(el)


def set_ea(rPr, face):
    el = rPr.find(qn("a:ea"))
    if el is None:
        el = etree.Element(qn("a:ea"))
        latin = rPr.find(qn("a:latin"))
        if latin is not None:
            latin.addnext(el)
        else:
            _insert_before(rPr, el, ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"))
    el.set("typeface", face)


def set_highlight(rPr, color):
    for old in rPr.findall(qn("a:highlight")):
        rPr.remove(old)
    hl = etree.Element(qn("a:highlight"))
    etree.SubElement(hl, qn("a:srgbClr"), val=color)
    _insert_before(rPr, hl, RPR_AFTER_HL)


def cjk_breaks(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("eaLnBrk", "1")
    pPr.set("hangingPunct", "1")
    pPr.set("latinLnBrk", "0")
    return pPr


def strip_style(shape):
    """Drop the theme style reference so no soft shadow / theme line sneaks in."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)
    return shape


# --------------------------------------------------------------------------- engine


class Canvas:
    def __init__(self, theme: dict, meta: dict, base_dir: Path):
        self.t, self.meta, self.base = theme, meta, base_dir
        c = theme["colors"]
        self.black, self.white, self.ground = c["black"], c["white"], c["ground"]
        self.muted, self.ink = c["muted"], c["text"]
        self.palette = c["palette"]
        self.chapter_cycle = c.get("chapter_cycle", self.palette)
        f = theme["fonts"]
        self.f_display, self.f_label, self.f_body = f["display"], f["label"], f["body"]
        self.f_ea, self.f_code = f["ea"], f.get("code", "Consolas")
        s = theme.get("stroke", {})
        self.bw, self.sw = s.get("border", 2.25), s.get("shadow", 5)
        self.tw, self.tsw = s.get("thin", 1.5), s.get("thin_shadow", 3)
        self.sizes = {"title": 28, "subtitle": 16, "body_max": 20, "body_min": 12, "split_below": 14,
                      **theme.get("sizes", {})}
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = E(W), E(H)
        self.blank = self.prs.slide_layouts[6]
        self.warnings: list[str] = []
        self.chapters = []
        self.counter = 0

    # ------------------------------------------------------------------ primitives
    def shape(self, s, kind, x, y, w, h, fill=None, line=None, lw=None, rot=0.0, pattern=None):
        sp = strip_style(s.shapes.add_shape(kind, E(x), E(y), E(w), E(h)))
        if pattern:
            sp.fill.patterned()
            sp.fill.pattern = pattern
            sp.fill.fore_color.rgb = rgb(fill[0])
            sp.fill.back_color.rgb = rgb(fill[1])
        elif fill:
            sp.fill.solid()
            sp.fill.fore_color.rgb = rgb(fill)
        else:
            sp.fill.background()
        if line:
            sp.line.color.rgb = rgb(line)
            sp.line.width = Pt(lw or self.bw)
        else:
            sp.line.fill.background()
        if rot:
            sp.rotation = rot
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = 0.5
        return sp

    def block(self, s, x, y, w, h, fill, thin=False, rot=0.0, kind=MSO_SHAPE.RECTANGLE, line=None,
              shadow=None):
        """Bordered shape with a hard, zero-blur offset shadow (BlockFrame's core atom)."""
        off = (self.tsw if thin else self.sw) / 72
        self.shape(s, kind, x + off, y + off, w, h, fill=shadow or self.black, rot=rot)
        return self.shape(s, kind, x, y, w, h, fill=fill, line=line or self.black,
                          lw=self.tw if thin else self.bw, rot=rot)

    def line(self, s, x1, y1, x2, y2, color=None, width=None):
        ln = strip_style(s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2)))
        ln.line.color.rgb = rgb(color or self.black)
        ln.line.width = Pt(width or self.bw)
        return ln

    def runs(self, p, text, size, font="body", bold=None, color=None, italic=None, spc=None, hl_color=None):
        latin = {"display": self.f_display, "label": self.f_label, "body": self.f_body,
                 "code": self.f_code}[font]
        for chunk, fmt in inline_runs(text):
            if not chunk:
                continue
            r = p.add_run()
            r.text = chunk
            f = r.font
            f.size = Pt(size)
            f.bold = True if fmt.get("bold") else bold
            if fmt.get("italic") or italic:
                f.italic = True
            f.color.rgb = rgb(color or self.ink)
            f.name = self.f_code if fmt.get("code") else latin
            rPr = r._r.get_or_add_rPr()
            set_ea(rPr, self.f_ea)
            if any(is_cjk(ch) for ch in chunk):
                rPr.set("lang", "zh-TW")
                rPr.set("altLang", "en-US")
            if spc:
                rPr.set("spc", str(int(spc * 100)))
            if fmt.get("hl"):
                set_highlight(rPr, (hl_color or self.palette[3]).lstrip("#"))
            if fmt.get("link"):
                r.hyperlink.address = fmt["link"]

    def tf_setup(self, tf, anchor=MSO_ANCHOR.TOP, margin=0.0, wrap=True):
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Emu(E(margin)))

    def text(self, s, x, y, w, h, lines, size, font="body", bold=None, color=None, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spc=None, line_spacing=None, target=None, margin=0.0):
        """lines: str or list[str|dict(text,size,bold,color,font,space_after)]."""
        if target is None:
            target = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
        tf = target.text_frame
        self.tf_setup(tf, anchor, margin)
        if isinstance(lines, str):
            lines = [lines]
        for k, ln in enumerate(lines):
            d = ln if isinstance(ln, dict) else {"text": ln}
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = d.get("align", align)
            cjk_breaks(p)
            if line_spacing:
                p.line_spacing = line_spacing
            if d.get("space_after") is not None:
                p.space_after = Pt(d["space_after"])
            self.runs(p, d["text"], d.get("size", size), d.get("font", font), d.get("bold", bold),
                      d.get("color", color), spc=d.get("spc", spc))
        return target

    def measure(self, text, size, font="body"):
        """Approximate rendered width in inches."""
        factor = {"display": 1.18, "label": 1.0, "body": 1.0, "code": 1.0}[font]
        w = 0.0
        for ch in plain(text):
            w += disp_width(ch) * (factor if ord(ch) < 0x2E80 else 1.0)
        return w * size / 72

    def pill(self, s, x, y, text, fill, size=11, color=None, rot=0.0, h=None, shadow=True, upper=True,
             font="label", line=None):
        text = text.upper() if upper and text.isascii() else text
        h = h or size / 72 * 2.0
        w = self.pill_width(text, size, font, upper)
        if shadow:
            sp = self.block(s, x, y, w, h, fill, thin=True, rot=rot, kind=MSO_SHAPE.ROUNDED_RECTANGLE, line=line)
        else:
            sp = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line or self.black,
                            lw=self.tw, rot=rot)
        self.text(s, 0, 0, 0, 0, text, size, font=font, bold=True, color=color or self.black,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spc=1 if upper else None, target=sp)
        return sp, w

    def lines_height(self, lines, width):
        """Estimated height (in) of text() lines inside `width` inches."""
        total = 0.0
        for ln in lines:
            d = ln if isinstance(ln, dict) else {"text": ln}
            sz = d.get("size", 14)
            n = max(1, math.ceil(self.measure(d["text"], sz) / max(0.5, width)))
            total += n * sz * 1.25 / 72 + d.get("space_after", 0) / 72
        return total

    def card_lines(self, it, title=20, desc=15, child=14):
        lines = [{"text": it["title"], "size": title, "bold": True, "color": self.black, "space_after": 3}]
        if it["desc"]:
            lines.append({"text": it["desc"], "size": desc, "color": "3A3A3A", "space_after": 2})
        for c in it["children"] + it["extra"]:
            lines.append({"text": "■ " + c, "size": child, "color": "3A3A3A"})
        return lines

    def pill_width(self, text, size, font="label", upper=True):
        return self.measure(text, size, font) + 0.36 + (len(text) * size * 0.08 / 72 if upper else 0)

    def star(self, s, x, y, d, fill, text=None, rot=12, size=14):
        sp = self.shape(s, MSO_SHAPE.STAR_10_POINT, x, y, d, d, fill=fill, line=self.black, lw=self.tw, rot=rot)
        if text:
            self.text(s, 0, 0, 0, 0, text, size, font="display", color=self.black, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE, target=sp)
        return sp

    def stripes(self, s, x, y, w, h, color, rot=0.0):
        return self.shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=(self.black, color), line=self.black,
                          lw=self.tw, rot=rot, pattern=MSO_PATTERN.WIDE_UPWARD_DIAGONAL)

    def dots(self, s, x, y, w, h):
        return self.shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=(self.muted, self.ground),
                          pattern=MSO_PATTERN.PERCENT_5)

    def picture(self, s, path, x, y, w, h):
        from PIL import Image
        p = Path(path)
        p = p if p.is_absolute() else (self.base / p)
        if not p.exists():
            self.warnings.append(f"找不到圖片：{path}")
            return None
        if p.suffix.lower() == ".svg":
            self.warnings.append(f"SVG 不支援，請先轉 PNG：{path}")
            return None
        with Image.open(p) as im:
            iw, ih = im.size
        sc = min(w / iw, h / ih)
        pw, ph = iw * sc, ih * sc
        pic = s.shapes.add_picture(str(p), E(x + (w - pw) / 2), E(y + (h - ph) / 2), E(pw), E(ph))
        pic._element.nvPicPr.cNvPr.set("descr", p.stem)
        return pic, (x + (w - pw) / 2, y + (h - ph) / 2, pw, ph)

    def color_for(self, k):
        return self.palette[k % len(self.palette)]

    def ch_color(self, ci):
        return self.chapter_cycle[ci % len(self.chapter_cycle)]

    def new_slide(self, ground=None, notes=None):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(ground or self.ground)
        self.counter += 1
        if notes:
            s.notes_slide.notes_text_frame.text = "\n\n".join(notes)
        return s

    # ------------------------------------------------------------------ structure
    def build(self, doc):
        meta = self.meta
        self.chapters = [c for c in doc.sections if c.title]
        pages = [Page("cover", title=doc.cover.title, subtitle=doc.cover.subtitle or meta.get("subtitle", ""),
                      notes=doc.cover.notes)]
        if doc.cover.blocks:
            pages.append(Page("content", title=doc.cover.title, blocks=doc.cover.blocks))
        if meta.get("outline", True) and len(self.chapters) >= 2:
            pages.append(Page("outline"))
        recap = bool(meta.get("recap", False))
        ci = -1
        for sec in doc.sections:
            if sec.title:
                ci += 1
                pages.append(Page("divider", chapter=ci, notes=sec.notes))
                intro = [b for b in sec.blocks]
                if intro and not (not sec.subsections and self.short_intro(intro)):
                    pages.append(Page("content", title=sec.title, chapter=ci, blocks=intro))
            elif sec.blocks:
                pages.append(Page("content", blocks=sec.blocks, notes=sec.notes))
            chap = ci if sec.title else None
            for sl in sec.slides:
                pages.append(Page("content", title=sl.title, subtitle=sl.subtitle, chapter=chap,
                                  blocks=sl.blocks, notes=sl.notes))
            for si, sub in enumerate(sec.subsections):
                if recap and si > 0 and chap is not None:
                    pages.append(Page("divider", chapter=chap, sub=si, recap=True))
                if sub.blocks or not sub.slides:
                    pages.append(Page("content", title=sub.title, subtitle=sub.subtitle, chapter=chap, sub=si,
                                      blocks=sub.blocks, notes=sub.notes))
                for sl in sub.slides:
                    pages.append(Page("content", title=sl.title, subtitle=sl.subtitle, chapter=chap, sub=si,
                                      blocks=sl.blocks, notes=sl.notes))
        closing = meta.get("closing", "Q&A")
        if closing:
            pages.append(Page("closing", title=str(closing) if closing is not True else "Q&A"))

        for pg in pages:
            if pg.kind == "content":
                for piece in self.paginate(pg):
                    self.render_content(piece)
            else:
                getattr(self, f"render_{pg.kind}")(pg)
        return self.prs

    @staticmethod
    def short_intro(blocks):
        return len(blocks) == 1 and blocks[0].kind in ("para", "quote") and disp_width(plain(blocks[0].data)) <= 80

    # ------------------------------------------------------------------ cover
    def render_cover(self, pg):
        s = self.new_slide(notes=pg.notes)
        m = self.meta
        self.dots(s, 7.4, 0, 5.93, 3.6)
        _, pw = self.pill(s, PAD + 0.1, 1.0, m.get("eyebrow", "PRESENTATION"), self.palette[3], size=13)
        title = pg.title
        # largest size (≤54pt) that splits the title into 1–3 evenly filled lines
        n = max(1.0, disp_width(title))
        size = max(min(54, int(7.1 * 72 / math.ceil(n / k) * 0.97)) for k in (1, 2, 3))
        lines = math.ceil(n * size / 72 / 7.1)
        tw = min(7.3, math.ceil(n / lines) * size / 72 * 1.04 + 0.15)   # narrow the box so lines balance
        self.text(s, PAD + 0.1, 1.75, tw, 3.2, title, size, bold=True, color=self.black,
                  anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
        if pg.subtitle:
            sub_font = "display" if pg.subtitle.isascii() else "body"
            sub = pg.subtitle.upper() if pg.subtitle.isascii() else pg.subtitle
            self.text(s, PAD + 0.1, 5.0, 7.3, 0.6, sub, 18, font=sub_font, bold=True, color=self.black)
        x = PAD + 0.1
        for k, key in enumerate(("author", "date")):
            if m.get(key):
                _, w = self.pill(s, x, 6.25, str(m[key]), self.white, size=12, upper=False, font="body")
                x += w + 0.25
        # right-hand decorations
        if m.get("cover_image"):
            r = self.picture_card(s, m["cover_image"], 8.3, 1.3, 4.3, 3.6, rot=3)
        else:
            self.block(s, 8.9, 1.2, 3.4, 2.5, self.palette[0], rot=6)
            self.block(s, 8.2, 3.3, 2.8, 2.1, self.palette[1], rot=-4)
            self.stripes(s, 11.1, 4.2, 1.5, 1.5, self.palette[2], rot=8)
        self.star(s, 11.6, 0.55, 1.2, self.palette[3], text=m.get("badge"), rot=14)

    def picture_card(self, s, path, x, y, w, h, rot=0.0):
        card = self.block(s, x, y, w, h, self.white, rot=rot)
        got = self.picture(s, path, x + 0.12, y + 0.12, w - 0.24, h - 0.24)
        if got and rot:
            got[0].rotation = rot
        return card

    # ------------------------------------------------------------------ outline (大綱)
    def render_outline(self, pg):
        s = self.new_slide()
        m = self.meta
        self.pill(s, PAD + 0.1, 1.0, m.get("outline_label", "OUTLINE"), self.palette[3], size=13)
        self.text(s, PAD + 0.1, 1.7, 5.8, 1.3, m.get("outline_title", "大綱"), 60, bold=True, color=self.black)
        self.text(s, PAD + 0.1, 2.9, 5.8, 0.7, m.get("outline_en", "OUTLINE"), 26, font="display",
                  color=self.black, spc=-0.5)
        self.block(s, 1.0, 4.6, 2.6, 1.8, self.palette[0], rot=-6)
        self.stripes(s, 3.3, 5.3, 1.3, 1.3, self.palette[1], rot=7)
        self.star(s, 4.5, 4.2, 1.0, self.palette[2], rot=10)

        n = len(self.chapters)
        top, bottom = 0.7, 6.8
        step = min(1.05, (bottom - top) / n)
        y0 = top + ((bottom - top) - step * n) / 2
        lx = 7.3
        self.line(s, lx, 0, lx, H, width=3)
        bh = min(0.62, step * 0.62)
        for k, ch in enumerate(self.chapters):
            cy = y0 + step * k + step / 2
            self.block(s, lx - 0.48, cy - bh / 2, 0.96, bh, self.ch_color(k))
            self.text(s, lx - 0.48, cy - bh / 2, 0.96, bh, f"{k + 1:02d}", 20, font="display",
                      color=self.black, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            lines = [{"text": ch.title, "size": 24, "bold": True, "color": self.black}]
            if ch.subtitle:
                lines.append({"text": ch.subtitle.upper(), "size": 11, "font": "label", "bold": True,
                              "color": self.muted, "spc": 1})
            self.text(s, lx + 0.9, cy - step / 2, 4.9, step, lines, 24, anchor=MSO_ANCHOR.MIDDLE)
        self.counter_pill(s)

    # ------------------------------------------------------------------ chapter divider + sub-TOC
    def render_divider(self, pg):
        ch = self.chapters[pg.chapter]
        col = self.ch_color(pg.chapter)
        s = self.new_slide(notes=pg.notes)
        pw = 5.3
        self.shape(s, MSO_SHAPE.RECTANGLE, -0.05, -0.05, pw + 0.05, H + 0.1, fill=col, line=self.black)
        self.text(s, PAD + 0.05, 0.45, 3.5, 1.4, f"{pg.chapter + 1:02d}", 80, font="display", color=self.black,
                  spc=-2)
        self.shape(s, MSO_SHAPE.RECTANGLE, PAD + 0.15, 1.85, 1.9, 0.1, fill=self.black)
        n = disp_width(ch.title)
        size = 60 if n <= 4 else 50 if n <= 6 else 40 if n <= 9 else 32
        lines = [{"text": ch.title, "size": size, "bold": True, "color": self.black, "space_after": 6}]
        if ch.subtitle:
            lines.append({"text": ch.subtitle.upper(), "size": 22, "font": "display", "color": self.black,
                          "spc": -0.3})
        self.text(s, 0.3, 2.5, pw - 0.6, 3.0, lines, size, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        self.star(s, pw - 0.75, 5.9, 1.2, self.palette[3] if col != self.palette[3] else self.palette[0], rot=16)

        subs = ch.subsections
        lx = 6.6
        if not subs:
            intro = ch.blocks[0].data if ch.blocks and self.short_intro(ch.blocks) else ""
            if intro:
                self.block(s, 6.4, 2.6, 6.2, 2.2, self.white)
                self.text(s, 6.7, 2.8, 5.6, 1.8, intro, 22, bold=True, color=self.black,
                          anchor=MSO_ANCHOR.MIDDLE)
            self.counter_pill(s)
            return
        self.line(s, lx, 0, lx, H, width=3)
        n = len(subs)
        top, bottom = 0.6, 6.9
        step = min(1.15, (bottom - top) / n)
        y0 = top + ((bottom - top) - step * n) / 2
        bh = min(0.66, step * 0.6)
        for k, sub in enumerate(subs):
            cy = y0 + step * k + step / 2
            current = pg.recap and k == pg.sub
            dim = pg.recap and k != pg.sub
            fill = self.black if current else (self.white if dim else col)
            self.block(s, lx - 0.45, cy - bh / 2, 0.9, bh, fill, thin=dim)
            self.text(s, lx - 0.45, cy - bh / 2, 0.9, bh, ROMAN[k] if k < len(ROMAN) else str(k + 1), 20,
                      font="display", color=self.white if current else (self.muted if dim else self.black),
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            lines = [{"text": sub.title, "size": 24, "bold": True,
                      "color": self.muted if dim else self.black}]
            if sub.subtitle:
                lines.append({"text": sub.subtitle, "size": 12, "color": self.muted})
            self.text(s, lx + 0.8, cy - step / 2, 5.6, step, lines, 24, anchor=MSO_ANCHOR.MIDDLE)
        self.counter_pill(s)

    # ------------------------------------------------------------------ closing
    def render_closing(self, pg):
        s = self.new_slide(ground=self.black)
        w, h = 8.2, 2.8
        x, y = (W - w) / 2, (H - h) / 2 + 0.2
        self.shape(s, MSO_SHAPE.RECTANGLE, x + 0.17, y + 0.17, w, h, fill=self.palette[3])
        frame = self.shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=self.black, line=self.white, lw=3)
        size = 72 if disp_width(pg.title) <= 6 else 48
        font = "display" if pg.title.isascii() else "body"
        self.text(s, 0, 0, 0, 0, pg.title, size, font=font, bold=True, color=self.white, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, target=frame)
        label = self.meta.get("closing_label", "THANK YOU")
        self.pill(s, (W - self.pill_width(label, 13)) / 2, y - 0.75, label, self.white, size=13, shadow=False)
        self.star(s, x + w - 0.7, y - 0.6, 1.3, self.palette[0], rot=18)

    # ------------------------------------------------------------------ chrome for content slides
    def counter_pill(self, s):
        self.pill(s, W - PAD - 0.75, H - 0.52, f"{self.counter:02d}", self.white, size=11, shadow=False)

    def nav(self, s, ci, si):
        col = self.ch_color(ci)
        self.shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, NAV_H, fill=col)
        self.line(s, 0, NAV_H, W, NAV_H)
        left, right = PAD, W - PAD
        if self.meta.get("logo_left"):
            self.picture(s, self.meta["logo_left"], PAD, 0.1, 0.8, 0.8)
            left += 1.0
        if self.meta.get("logo_right"):
            self.picture(s, self.meta["logo_right"], W - PAD - 0.8, 0.1, 0.8, 0.8)
            right -= 1.0
        n = len(self.chapters)
        slot = (right - left) / n
        for k, ch in enumerate(self.chapters):
            cx = left + slot * k + slot / 2
            if k == ci:
                tw = self.measure(ch.title, 14) + 0.5
                sp = self.block(s, cx - tw / 2, 0.1, tw, 0.4, self.white, thin=True)
                self.text(s, 0, 0, 0, 0, ch.title, 14, bold=True, color=self.black, align=PP_ALIGN.CENTER,
                          anchor=MSO_ANCHOR.MIDDLE, target=sp)
            else:
                self.text(s, cx - slot / 2, 0.1, slot, 0.4, ch.title, 13, color="4A4A4A",
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        subs = self.chapters[ci].subsections
        if not subs:
            return
        size = 11
        widths = [self.measure(sub.title, size) + 0.4 for sub in subs]
        gap = 0.18
        total = sum(widths) + gap * (len(subs) - 1)
        if total > right - left:
            size = 9
            widths = [self.measure(sub.title, size) + 0.3 for sub in subs]
            total = sum(widths) + gap * (len(subs) - 1)
        cx = left + slot * ci + slot / 2
        x = min(max(left, cx - total / 2), right - total)
        for k, (sub, w) in enumerate(zip(subs, widths)):
            if k == si:
                sp = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 0.6, w, 0.3, fill=self.black, line=self.black,
                                lw=self.tw)
                self.text(s, 0, 0, 0, 0, sub.title, size, bold=True, color=self.white, align=PP_ALIGN.CENTER,
                          anchor=MSO_ANCHOR.MIDDLE, target=sp)
            else:
                self.text(s, x, 0.6, w, 0.3, sub.title, size, color="4A4A4A", align=PP_ALIGN.CENTER,
                          anchor=MSO_ANCHOR.MIDDLE)
            x += w + gap

    def full_title(self, pg):
        if pg.chapter is None or pg.sub is None:
            return pg.title
        sub = self.chapters[pg.chapter].subsections[pg.sub].title
        if pg.title == sub or pg.title.startswith(sub):
            return pg.title
        fmt = self.meta.get("title_format", "{sub} – {title}")
        return fmt.format(sub=sub, title=pg.title)

    # ------------------------------------------------------------------ pagination
    def split_blocks(self, blocks):
        text = [b for b in blocks if b.kind in TEXT_KINDS]
        visuals = [b for b in blocks if b.kind in ("image", "table", "code", "component")]
        callouts = [b for b in blocks if b.kind == "callout"]
        sources = [b for b in blocks if b.kind == "source"]
        return text, visuals, callouts, sources

    def body_box(self, has_sub, n_callouts):
        top = NAV_H + 1.35 + (0.4 if has_sub else 0)
        bottom = H - 0.75 - (1.15 * n_callouts)
        return PAD, top, W - 2 * PAD, bottom - top

    def text_ratio(self, visuals):
        if any(v.kind == "component" and v.data["type"] in ("cards", "steps", "compare") for v in visuals):
            return 0.42
        return 0.45

    def paginate(self, pg):
        text, visuals, callouts, sources = self.split_blocks(pg.blocks)
        max_rows = int(self.meta.get("max_table_rows", 9))
        vis = []
        for v in visuals:
            if v.kind == "table" and len(v.data["rows"]) - 1 > max_rows:
                head, body = v.data["rows"][0], v.data["rows"][1:]
                for k in range(0, len(body), max_rows):
                    vis.append(Block("table", {"rows": [head] + body[k:k + max_rows], "aligns": v.data["aligns"]}))
            else:
                vis.append(v)
        first_vis, rest_vis = (vis[:1], vis[1:]) if text else (vis[:2], vis[2:])
        if len(first_vis) == 2 and not all(v.kind == "image" for v in first_vis):
            rest_vis = first_vis[1:] + rest_vis
            first_vis = first_vis[:1]
        pages = self.fit_text(pg, text, first_vis, callouts, sources)
        for v in rest_vis:
            pages.append(Page("content", title=pg.title, subtitle=pg.subtitle, chapter=pg.chapter, sub=pg.sub,
                              blocks=[v]))
        suffix = self.meta.get("continued_suffix", "（續）")
        for p in pages[1:]:
            if not p.title.endswith(suffix):
                p.title = pg.title + suffix
        return pages

    def fit_text(self, pg, text, vis, callouts, sources):
        base = Page("content", title=pg.title, subtitle=pg.subtitle, chapter=pg.chapter, sub=pg.sub,
                    notes=pg.notes)
        units = [u for b in text for u, _ in text_units(b)]
        _, _, w, h = self.body_box(bool(pg.subtitle), len(callouts))
        if vis:
            w *= self.text_ratio(vis)
        floor = self.sizes["split_below"]
        paras = self.text_paras(text)
        if len(units) < 2 or estimate_height_pt(paras, floor * 1.08, E(w)) <= h * 72:
            base.blocks = text + vis + callouts + sources
            return [base]
        mid = len(units) // 2
        a = self.fit_text(pg, merge_bullets(units[:mid]), vis, callouts, sources)
        b_pg = Page("content", title=pg.title, subtitle=pg.subtitle, chapter=pg.chapter, sub=pg.sub)
        b = self.fit_text(b_pg, merge_bullets(units[mid:]), [], [], [])
        return a + b

    # ------------------------------------------------------------------ content slide
    def text_paras(self, blocks):
        paras = []
        for b in blocks:
            if b.kind == "bullets":
                for it in b.data:
                    paras.append({"text": it["text"], "level": it["level"], "ordered": it["ordered"],
                                  "bullet": "sq", "scale": max(0.8, 1 - 0.1 * it["level"])})
            elif b.kind == "para":
                paras.append({"text": b.data, "bullet": "none"})
            elif b.kind == "sub":
                paras.append({"text": b.data, "bullet": "none", "bold": True, "scale": 1.05})
            elif b.kind == "quote":
                paras.append({"text": f"「{b.data}」", "bullet": "none", "italic": True, "color": self.muted})
        return paras

    def render_content(self, pg):
        s = self.new_slide(notes=pg.notes)
        text, visuals, callouts, sources = self.split_blocks(pg.blocks)
        if pg.chapter is not None:
            self.nav(s, pg.chapter, pg.sub)
        top = NAV_H + 0.3 if pg.chapter is not None else 0.5
        self.text(s, PAD, top, W - 2 * PAD, 0.7, self.full_title(pg), self.sizes["title"], bold=True,
                  color=self.black, anchor=MSO_ANCHOR.MIDDLE)
        if pg.subtitle:
            self.text(s, PAD, top + 0.7, W - 2 * PAD, 0.4, pg.subtitle, self.sizes["subtitle"], bold=True,
                      color=self.muted)
        x, y, w, h = self.body_box(bool(pg.subtitle), len(callouts))
        if pg.chapter is None:
            y -= NAV_H - 0.2
            h += NAV_H - 0.2
        paras = self.text_paras(text)

        statement = (not visuals and paras and len(paras) <= 4 and all(p["bullet"] == "none" for p in paras)
                     and any(f.get("hl") for p in paras for _, f in inline_runs(p["text"])))
        only_quote = not visuals and len(text) == 1 and text[0].kind == "quote"
        if statement:
            self.draw_statement(s, paras, x, y, w, h)
        elif only_quote:
            self.draw_quote(s, text[0].data, x, y, w, h)
        elif not visuals:
            self.draw_text(s, paras, x, y, w, h, pg.title, max_size=self.sizes["body_max"] + 4)
        elif not paras:
            self.draw_visuals(s, visuals, x, y, w, h, pg.title)
        else:
            tw = w * self.text_ratio(visuals)
            gap = 0.45
            self.draw_text(s, paras, x, y, tw, h, pg.title)
            self.draw_visuals(s, visuals, x + tw + gap, y, w - tw - gap, h, pg.title)

        cy = y + h + 0.3
        for c in callouts:
            self.draw_callout(s, c.data, x, cy, w, 0.85)
            cy += 1.15
        if sources:
            src = "；".join(plain(b.data) for b in sources)
            self.text(s, PAD, H - 0.55, W - 2 * PAD - 1.2, 0.4, src, 9, color=self.muted,
                      anchor=MSO_ANCHOR.MIDDLE)
        self.counter_pill(s)

    def draw_text(self, s, paras, x, y, w, h, title, max_size=None):
        if not paras:
            return
        size = self.sizes["body_min"]
        for sz in range(int(max_size or self.sizes["body_max"]), int(self.sizes["body_min"]) - 1, -1):
            if estimate_height_pt(paras, sz * 1.08, E(w)) <= h * 72:
                size = sz
                break
        else:
            self.warnings.append(f"「{title}」文字可能溢出，建議拆頁或精簡")
        tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
        tf = tb.text_frame
        self.tf_setup(tf)
        num = {}
        for k, pd in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            lvl = pd.get("level", 0)
            sz = max(10, round(size * pd.get("scale", 1.0)))
            p.space_after = Pt(round(sz * 0.55))
            p.line_spacing = 1.12
            pPr = cjk_breaks(p)
            if pd["bullet"] == "sq":
                ind = 0.32 + 0.3 * lvl
                pPr.set("marL", str(E(ind)))
                pPr.set("indent", str(-E(0.28)))
                if pd.get("ordered"):
                    num[lvl] = num.get(lvl, 0) + 1
                    bu = etree.SubElement(pPr, qn("a:buFont"), typeface=self.f_display)
                    bu = etree.SubElement(pPr, qn("a:buAutoNum"), type="arabicPeriod", startAt=str(num[lvl]))
                else:
                    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
                    etree.SubElement(pPr, qn("a:buChar"), char="■" if lvl == 0 else "–")
            else:
                num.clear()
                pPr.set("marL", "0")
                pPr.set("indent", "0")
                etree.SubElement(pPr, qn("a:buNone"))
            self.runs(p, pd["text"], sz, bold=pd.get("bold"), color=pd.get("color", self.ink),
                      italic=pd.get("italic"))

    def draw_statement(self, s, paras, x, y, w, h):
        lines = [{"text": p["text"], "size": 28, "bold": True, "space_after": 22} for p in paras]
        self.text(s, x + 0.6, y, w - 1.2, h, lines, 28, color=self.black, anchor=MSO_ANCHOR.MIDDLE,
                  align=PP_ALIGN.CENTER)
        self.star(s, x + w - 1.1, y + 0.1, 0.9, self.palette[0], rot=12)

    def draw_quote(self, s, text, x, y, w, h):
        cw, ch = min(w, 10.0), min(h, 3.2)
        cx, cy = x + (w - cw) / 2, y + (h - ch) / 2
        self.block(s, cx, cy, cw, ch, self.white)
        self.text(s, cx + 0.5, cy + 0.3, cw - 1.0, ch - 0.6, f"「{text}」", 30, bold=True, color=self.black,
                  anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        self.pill(s, cx - 0.2, cy - 0.3, "QUOTE", self.palette[0], size=12, rot=-5)

    CALLOUT = {"note": ("註", 1), "tip": ("TIP", 2), "important": ("重點", 0), "warning": ("注意", 3),
               "caution": ("注意", 0), "summary": ("小結", 2)}

    def draw_callout(self, s, d, x, y, w, h):
        label, ci = self.CALLOUT.get(d["kind"], (d["kind"].upper(), 1))
        self.block(s, x, y, w, h, self.color_for(ci), thin=True)
        self.text(s, x + 1.3, y, w - 1.6, h, d["text"], 15, bold=True, color=self.black, anchor=MSO_ANCHOR.MIDDLE)
        self.pill(s, x + 0.2, y + (h - 0.44) / 2, label, self.white, size=13, rot=-4)

    # ------------------------------------------------------------------ visuals
    def draw_visuals(self, s, visuals, x, y, w, h, title):
        n = len(visuals)
        gap = 0.4
        if n == 1:
            cells = [(x, y, w, h)]
        else:
            cw = (w - gap * (n - 1)) / n
            cells = [(x + k * (cw + gap), y, cw, h) for k in range(n)]
        for v, cell in zip(visuals, cells):
            if v.kind == "component":
                getattr(self, f"comp_{v.data['type'].replace('timeline', 'steps')}")(s, v.data, *cell)
            else:
                getattr(self, f"draw_{v.kind}")(s, v.data, *cell, title=title)

    def draw_image(self, s, d, x, y, w, h, title=""):
        from PIL import Image
        p = Path(d["path"])
        p = p if p.is_absolute() else self.base / p
        if not p.exists() or p.suffix.lower() == ".svg":
            self.warnings.append(f"「{title}」圖片無法使用：{d['path']}")
            return
        cap = 0.55 if d.get("alt") else 0.1
        with Image.open(p) as im:
            iw, ih = im.size
        pad = 0.12
        sc = min((w - 2 * pad - 0.1) / iw, (h - cap - 2 * pad - 0.1) / ih)
        fw, fh = iw * sc + 2 * pad, ih * sc + 2 * pad
        fx, fy = x + (w - fw) / 2, y + (h - cap - fh) / 2
        self.block(s, fx, fy, fw, fh, self.white)
        s.shapes.add_picture(str(p), E(fx + pad), E(fy + pad), E(fw - 2 * pad), E(fh - 2 * pad))
        if d.get("alt"):
            self.pill(s, fx + 0.2, fy + fh - 0.12, d["alt"], self.palette[3], size=11, rot=-3, upper=False,
                      font="body")

    def draw_code(self, s, d, x, y, w, h, title=""):
        lines = d["text"].split("\n")
        size = 16
        longest = max((len(l) for l in lines), default=1)
        while size > 9 and (longest * size * 0.6 / 72 > w - 0.6 or len(lines) * size * 1.3 / 72 > h - 0.6):
            size -= 1
        ch = min(h, len(lines) * size * 1.3 / 72 + 0.7)
        cy = y + (h - ch) / 2
        self.block(s, x, cy, w, ch, "1A1A1A")
        tb = self.text(s, x + 0.3, cy + 0.35, w - 0.6, ch - 0.5,
                       [{"text": l or " ", "font": "code"} for l in lines], size, font="code", color="F5F5F5")
        self.pill(s, x + 0.2, cy - 0.2, d.get("lang") or "CODE", self.palette[3], size=11, rot=-3)

    def draw_table(self, s, d, x, y, w, h, title="", header_color=None):
        rows = d["rows"]
        nc = max(len(r) for r in rows)
        rows = [r + [""] * (nc - len(r)) for r in rows]
        nr = len(rows)
        size = max(10, min(16, int(22 - 0.8 * nr - 0.6 * nc)))
        rh = size * 2.2 / 72
        th = min(h, rh * nr)
        tx = x
        gf = s.shapes.add_table(nr, nc, E(tx), E(y), E(w), E(th))
        # shadow must sit behind the table
        shadow = self.shape(s, MSO_SHAPE.RECTANGLE, tx + self.sw / 72, y + self.sw / 72, w, th, fill=self.black)
        gf._element.addprevious(shadow._element)
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        tblPr.set("bandRow", "0")
        tblPr.set("firstRow", "1")
        widths = [max(disp_width(plain(r[c])) for r in rows) + 2 for c in range(nc)]
        tot = sum(widths)
        for c in range(nc):
            tbl.columns[c].width = E(w * widths[c] / tot)
        aligns = d.get("aligns") or []
        amap = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
        hc = header_color or self.palette[2]
        for r, row in enumerate(rows):
            tbl.rows[r].height = E(rh)
            for c, val in enumerate(row):
                cell = tbl.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(hc if r == 0 else (self.white if r % 2 else self.ground))
                cell.margin_left = cell.margin_right = E(0.08)
                cell.margin_top = cell.margin_bottom = E(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tcPr = cell._tc.get_or_add_tcPr()
                for k, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
                    ln = etree.Element(qn(tag), w=str(int(self.tw * 12700)), cap="flat", cmpd="sng", algn="ctr")
                    sf = etree.SubElement(ln, qn("a:solidFill"))
                    etree.SubElement(sf, qn("a:srgbClr"), val=self.black)
                    tcPr.insert(k, ln)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                v = val.strip()
                if c < len(aligns) and aligns[c] != "l":
                    p.alignment = amap[aligns[c]]
                elif r == 0 or re.fullmatch(r"[○●◎×✓✗✔✘△\-—vVxXoO]|[\d.,%+\-]+", v or "-"):
                    p.alignment = PP_ALIGN.CENTER
                self.runs(p, val, size, bold=True if r == 0 else None, color=self.black)
        if rh * nr > h * 1.05:
            self.warnings.append(f"「{title}」表格可能超出版面（{nr} 列）")

    # ------------------------------------------------------------------ components
    def item_icon(self, s, x, y, d, k, it, size=18):
        self.block(s, x, y, d, d, self.color_for(k + 1), thin=True)
        label = it["attrs"].get("icon") or f"{k + 1}"
        font = "display" if label.isascii() else "body"
        self.text(s, x, y, d, d, label, size, font=font, color=self.black, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE)

    def comp_cards(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        args = d["args"]
        grid = "grid" in args or ("stack" not in args and w > 8)
        gap = 0.35
        if not grid:
            ch = min(1.6, (h - gap * (n - 1)) / n)
            y0 = y + (h - (ch * n + gap * (n - 1))) / 2
            for k, it in enumerate(items):
                self.card(s, it, k, x + 0.25, y0 + k * (ch + gap), w - 0.35, ch, horizontal=True)
        else:
            cols = int(next((a.split("=")[1] for a in args if a.startswith("cols=")), 0)) or (
                n if n <= 3 else 2 if n == 4 else 3)
            rows = -(-n // cols)
            cw = (w - gap * (cols - 1)) / cols
            need = max(self.lines_height(self.card_lines(it), cw - 0.6) for it in items) + 0.62 + 0.9
            ch = min(max(need, 2.6), (h - gap * (rows - 1)) / rows)
            y0 = y + (h - (ch * rows + gap * (rows - 1))) / 2
            for k, it in enumerate(items):
                r, c = divmod(k, cols)
                self.card(s, it, k, x + c * (cw + gap), y0 + r * (ch + gap), cw - 0.1, ch, horizontal=False)

    def card(self, s, it, k, x, y, w, h, horizontal):
        self.block(s, x, y, w, h, self.white)
        img = it["attrs"].get("img")
        inner_x, inner_w = x + 0.25, w - 0.5
        if img:
            iw = min(h - 0.3, w * 0.3)
            self.picture(s, img, x + w - iw - 0.18, y + (h - iw) / 2, iw, iw)
            inner_w -= iw + 0.15
        lines = self.card_lines(it)
        if horizontal:
            d = min(0.62, h - 0.4)
            self.item_icon(s, inner_x, y + (h - d) / 2, d, k, it)
            self.text(s, inner_x + d + 0.25, y + 0.1, inner_w - d - 0.25, h - 0.2, lines, 20,
                      anchor=MSO_ANCHOR.MIDDLE)
        else:
            d = 0.62
            self.item_icon(s, inner_x, y + 0.3, d, k, it)
            self.text(s, inner_x, y + 0.3 + d + 0.25, inner_w, h - d - 0.7, lines, 20)
        if it["attrs"].get("tag"):
            self.pill(s, x - 0.3, y - 0.28, it["attrs"]["tag"], self.color_for(k + 3), size=12, rot=-8,
                      upper=False, font="body")

    def comp_steps(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        horizontal = d["type"] == "timeline" or "horizontal" in d["args"] or (w > 8 and n <= 5
                                                                                and "vertical" not in d["args"])
        if horizontal:
            gap = 0.55
            cw = (w - gap * (n - 1)) / n
            need = 0.55 + 0.4 + max(self.lines_height(
                [{"text": it["title"], "size": 20, "space_after": 6}, {"text": it["desc"] or "", "size": 15}]
                + [{"text": "■ " + c, "size": 14} for c in it["children"]], cw - 0.4) for it in items)
            ch = min(h, max(2.8, need))
            cy = y + (h - ch) / 2
            for k in range(n - 1):
                x1 = x + (k + 1) * cw + k * gap
                self.line(s, x1, cy + ch / 2, x1 + gap, cy + ch / 2, width=4)
            for k, it in enumerate(items):
                cx = x + k * (cw + gap)
                self.block(s, cx, cy, cw, ch, self.white, thin=True)
                self.shape(s, MSO_SHAPE.RECTANGLE, cx, cy, cw, 0.55, fill=self.color_for(k), line=self.black,
                           lw=self.tw)
                self.text(s, cx, cy, cw, 0.55, f"STEP {k + 1}", 15, font="label", bold=True, color=self.black,
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spc=1.5)
                lines = [{"text": it["title"], "size": 20, "bold": True, "color": self.black, "space_after": 6}]
                if it["desc"]:
                    lines.append({"text": it["desc"], "size": 15, "color": "3A3A3A"})
                lines += [{"text": "■ " + c, "size": 14, "color": "3A3A3A"} for c in it["children"]]
                self.text(s, cx + 0.2, cy + 0.75, cw - 0.4, ch - 0.9, lines, 20)
            return
        step = min(1.35, h / n)
        y0 = y + (h - step * n) / 2
        dsz = min(0.6, step * 0.62)
        lx = x + 0.2 + dsz / 2
        if n > 1:
            self.line(s, lx, y0 + step / 2, lx, y0 + step * (n - 0.5), width=3)
        for k, it in enumerate(items):
            cy = y0 + k * step + step / 2
            self.block(s, lx - dsz / 2, cy - dsz / 2, dsz, dsz, self.color_for(k), thin=True)
            self.text(s, lx - dsz / 2, cy - dsz / 2, dsz, dsz, str(k + 1), 18, font="display", color=self.black,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            lines = [{"text": it["title"], "size": 20, "bold": True, "color": self.black, "space_after": 2}]
            if it["desc"]:
                lines.append({"text": it["desc"], "size": 15, "color": "3A3A3A"})
            lines += [{"text": "■ " + c, "size": 14, "color": "3A3A3A"} for c in it["children"]]
            self.text(s, lx + dsz / 2 + 0.3, cy - step / 2, w - dsz - 0.7, step, lines, 20,
                      anchor=MSO_ANCHOR.MIDDLE)

    def comp_flow(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        vertical = "vertical" in d["args"] or (w / h < 1.3 and "horizontal" not in d["args"])
        aw = 0.42
        if vertical:
            bh = min(1.25, (h - aw * (n - 1)) / n)
            bw = min(w, 4.6)
            bx = x + (w - bw) / 2
            y0 = y + (h - (bh * n + aw * (n - 1))) / 2
            for k, it in enumerate(items):
                by = y0 + k * (bh + aw)
                self.flow_box(s, it, k, bx, by, bw, bh)
                if k < n - 1:
                    self.shape(s, MSO_SHAPE.DOWN_ARROW, bx + bw / 2 - 0.17, by + bh + 0.06, 0.34, aw - 0.12,
                               fill=self.black)
            return
        per_row = n if n <= 6 else -(-n // 2)
        rows = -(-n // per_row)
        bw = (w - aw * (per_row - 1)) / per_row
        bh = min(2.2, (h - 0.5 * (rows - 1)) / rows)
        y0 = y + (h - (bh * rows + 0.5 * (rows - 1))) / 2
        for k, it in enumerate(items):
            r, c = divmod(k, per_row)
            bx, by = x + c * (bw + aw), y0 + r * (bh + 0.5)
            self.flow_box(s, it, k, bx, by, bw - 0.05, bh)
            if c < per_row - 1 and k < n - 1:
                self.shape(s, MSO_SHAPE.RIGHT_ARROW, bx + bw + 0.02, by + bh / 2 - 0.16, aw - 0.12, 0.32,
                           fill=self.black)

    def flow_box(self, s, it, k, x, y, w, h):
        self.block(s, x, y, w, h, self.color_for(k), thin=True)
        lines = [{"text": it["title"], "size": 19, "bold": True, "color": self.black, "space_after": 3}]
        if it["desc"]:
            lines.append({"text": it["desc"], "size": 14, "color": "2A2A2A"})
        self.text(s, x + 0.12, y + 0.05, w - 0.24, h - 0.1, lines, 19, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE)

    def comp_stats(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        cols = n if (w > 8 or n <= 2) else 2
        rows = -(-n // cols)
        gap = 0.45
        cw = (w - gap * (cols - 1)) / cols
        ch = min(3.0, (h - gap * (rows - 1)) / rows)
        y0 = y + (h - (ch * rows + gap * (rows - 1))) / 2
        for k, it in enumerate(items):
            r, c = divmod(k, cols)
            cx, cy = x + c * (cw + gap), y0 + r * (ch + gap)
            rot = -2 if k % 2 == 0 else 2
            self.block(s, cx, cy, cw, ch, self.white, thin=True, rot=rot)
            self.shape(s, MSO_SHAPE.OVAL, cx + cw - 0.42, cy + 0.2, 0.2, 0.2, fill=self.color_for(k), line=self.black,
                       lw=self.tw)
            num = it["title"]
            size = 60 if len(num) <= 5 else 44 if len(num) <= 8 else 32
            lines = [{"text": num, "size": size, "font": "display", "color": self.black, "space_after": 4}]
            if it["desc"]:
                lines.append({"text": it["desc"], "size": 20, "bold": True, "color": self.black, "space_after": 2})
            for extra in it["extra"] + it["children"]:
                lines.append({"text": extra, "size": 14, "color": "3A3A3A"})
            tb = self.text(s, cx + 0.3, cy + 0.2, cw - 0.6, ch - 0.4, lines, size, anchor=MSO_ANCHOR.MIDDLE)
            tb.rotation = rot

    def comp_compare(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        gap = 0.6 if n == 2 else 0.4
        cw = (w - gap * (n - 1)) / n
        hh = 0.95
        need = hh + 0.7 + max(self.lines_height(
            [{"text": "■ " + c, "size": 19, "space_after": 10} for c in it["children"] + it["extra"]], cw - 0.7)
            for it in items)
        ch = min(h, max(2.6, need))
        y += (h - ch) / 2
        h = ch
        for k, it in enumerate(items):
            cx = x + k * (cw + gap)
            self.block(s, cx, y, cw, h, self.white)
            self.shape(s, MSO_SHAPE.RECTANGLE, cx, y, cw, hh, fill=self.color_for(k * 2), line=self.black,
                       lw=self.bw)
            head = [{"text": it["title"], "size": 24, "bold": True, "color": self.black}]
            if it["desc"]:
                head.append({"text": it["desc"], "size": 13, "color": "2A2A2A"})
            self.text(s, cx, y, cw, hh, head, 24, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            lines = [{"text": "■ " + c, "size": 19, "color": self.ink, "space_after": 10}
                     for c in it["children"] + it["extra"]]
            if lines:
                self.text(s, cx + 0.35, y + hh + 0.35, cw - 0.7, h - hh - 0.5, lines, 19)
        if n == 2:
            self.star(s, x + cw + gap / 2 - 0.55, y + h / 2 - 0.55, 1.1, self.palette[3], text="VS", rot=10, size=16)

    def comp_chart(self, s, d, x, y, w, h):
        rows = d["rows"]
        if len(rows) < 2:
            self.warnings.append("chart 區塊需要一個 Markdown 表格（第一列為表頭）")
            return
        kind = (d["args"][0] if d["args"] else "column").lower()
        title = " ".join(d["args"][1:])
        types = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                 "stacked": XL_CHART_TYPE.COLUMN_STACKED, "line": XL_CHART_TYPE.LINE_MARKERS,
                 "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT}
        ct = types.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
        cd = CategoryChartData()
        cats = [r[0] for r in rows[1:]]
        cd.categories = cats

        def num(v):
            try:
                return float(re.sub(r"[,%\s]", "", v))
            except ValueError:
                return 0.0

        series_names = rows[0][1:]
        if kind in ("pie", "doughnut"):
            series_names = series_names[:1]
        for j, name in enumerate(series_names):
            cd.add_series(name, [num(r[j + 1]) if j + 1 < len(r) else 0 for r in rows[1:]])
        self.block(s, x, y, w, h, self.white)
        top = 0.25
        if title:
            self.pill(s, x + 0.25, y - 0.2, title, self.palette[3], size=12, rot=-2, upper=False, font="body")
            top = 0.4
        gf = s.shapes.add_chart(ct, E(x + 0.2), E(y + top), E(w - 0.4), E(h - top - 0.15), cd)
        ch = gf.chart
        ch.font.size = Pt(12)
        ch.font.name = self.f_body
        set_ea(ch.font._rPr, self.f_ea)
        multi = len(series_names) > 1 or kind in ("pie", "doughnut")
        ch.has_legend = multi
        if multi:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        plot = ch.plots[0]
        if kind in ("pie", "doughnut"):
            ser = plot.series[0]
            for i in range(len(cats)):
                pt = ser.points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(self.color_for(i))
                pt.format.line.color.rgb = rgb(self.black)
                pt.format.line.width = Pt(1.5)
        else:
            if kind in ("column", "bar", "stacked"):
                plot.gap_width = 60
                if kind == "stacked":
                    plot.overlap = 100
            for i, ser in enumerate(plot.series):
                c = self.color_for(i + 2 if len(plot.series) == 1 else i)
                if kind == "line":
                    ser.format.line.color.rgb = rgb(self.black if len(plot.series) == 1 else c)
                    ser.format.line.width = Pt(3)
                    ser.smooth = False
                    ser.marker.format.fill.solid()
                    ser.marker.format.fill.fore_color.rgb = rgb(c)
                    ser.marker.format.line.color.rgb = rgb(self.black)
                    ser.marker.size = 9
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = rgb(c)
                    ser.format.line.color.rgb = rgb(self.black)
                    ser.format.line.width = Pt(1.5)
            if len(cats) <= 8 and len(plot.series) <= 2:
                plot.has_data_labels = True
                plot.data_labels.font.size = Pt(11)
                plot.data_labels.font.bold = True
            va = ch.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = rgb("DDDDDD")
            va.format.line.fill.background()
            ch.category_axis.format.line.color.rgb = rgb(self.black)
            ch.category_axis.format.line.width = Pt(2)


def build_canvas(doc, theme, base_dir: Path):
    c = Canvas(theme, doc.meta, base_dir)
    prs = c.build(doc)
    return prs, c
